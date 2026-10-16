#!/usr/bin/env python3
"""
BotBhai lead finder — free, legal lead lists of local businesses.

Source: OpenStreetMap via the public Overpass API (data © OpenStreetMap contributors, ODbL licence —
keep the attribution line in every export). No API key, no cost, no scraping of sites that forbid it
(Google Maps and Justdial terms prohibit scraping, so we don't touch them).

Usage:
  python leadfinder.py --city Hyderabad --niche salon gym realestate examkits --out leads
Outputs: leads.csv and leads.xlsx (one sheet per niche + a summary), scored for BotBhai outreach.
Runs free in the cloud via GitHub Actions (.github/workflows/leadfinder.yml) — nothing on your PC.
"""
import argparse, csv, json, sys, time, urllib.parse, urllib.request

ENDPOINTS = ["https://overpass-api.de/api/interpreter", "https://overpass.kumi.systems/api/interpreter"]
NICHE_TAGS = {
    "salon":      ['shop~"^(hairdresser|beauty|cosmetics)$"', 'amenity="spa"'],
    "gym":        ['leisure~"^(fitness_centre|sports_centre)$"', 'sport~"^(fitness|yoga|crossfit)$"'],
    "realestate": ['office~"^(estate_agent|property_management)$"', 'shop="estate_agent"'],
    "examkits":   ['amenity~"^(college|training|prep_school)$"', 'office="educational_institution"'],
}
PITCH = {
    "salon": "Clients book slots and see offers on WhatsApp; fill empty afternoon chairs.",
    "gym": "Free-trial bookings and renewal reminders on WhatsApp; recover lapsed members.",
    "realestate": "Qualify buyers (budget, BHK, area) and book site visits automatically, even at night.",
    "examkits": "Daily quiz on WhatsApp builds a list of aspirants who later buy courses.",
}

def overpass(query, retries=3):
    data = urllib.parse.urlencode({"data": query}).encode()
    last = None
    for attempt in range(retries):
        for url in ENDPOINTS:
            try:
                req = urllib.request.Request(url, data=data, headers={"User-Agent": "KalpaLabs-BotBhai-LeadFinder/1.0 (contact via github.com/Akhilkanigicharla)"})
                with urllib.request.urlopen(req, timeout=180) as r:
                    return json.load(r)
            except Exception as e:  # rate limit / busy server → wait and try the mirror
                last = e; time.sleep(10 * (attempt + 1))
    raise SystemExit(f"Overpass failed: {last}")

def build_query(city, tags, near=None, radius=1000):
    """Whole city by default; with near=(lat, lon) only a walking cluster (Overpass "around" filter, radius in metres)."""
    if near:
        parts = "".join(f'nwr(around:{int(radius)},{near[0]},{near[1]})[{t}];' for t in tags)
        return f'[out:json][timeout:60];({parts});out center tags;'
    parts = "".join(f'nwr[{t}](area.city);' for t in tags)
    return f'[out:json][timeout:170];area["name"="{city}"]["boundary"="administrative"]->.city;({parts});out center tags;'

def parse_near(text):
    try:
        lat, lon = (float(x) for x in str(text).split(","))
    except ValueError:
        raise SystemExit(f"--near must be 'lat,lon', got {text!r}")
    if not (6 <= lat <= 37 and 68 <= lon <= 98):
        raise SystemExit(f"--near {text!r} is outside India")
    return lat, lon

def is_late(hours):
    """True if any opening range ends at or after 9 PM (or runs past midnight): their WhatsApps keep coming late."""
    import re
    for a, b in re.findall(r"(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})", hours or ""):
        if b >= "21:00" or b < a: return True
    return False

def phone_of(t):
    for k in ("contact:phone", "phone", "contact:mobile", "mobile", "contact:whatsapp"):
        if t.get(k): return t[k].split(";")[0].strip()
    return ""

def score(row):
    """Higher = better BotBhai prospect. Has a phone (we can reach them) but weak online presence (we can help)."""
    s = 0
    if row["phone"]: s += 40
    if not row["website"]: s += 25
    if not row["opening_hours"]: s += 10
    if row["address"]: s += 10
    if row["email"]: s += 5
    if row["name"]: s += 10
    return s

def collect(city, niche, near=None, radius=1000):
    res = overpass(build_query(city, NICHE_TAGS[niche], near, radius))
    seen, rows = set(), []
    for el in res.get("elements", []):
        t = el.get("tags", {}); name = t.get("name", "").strip()
        if not name: continue
        lat = el.get("lat") or el.get("center", {}).get("lat"); lon = el.get("lon") or el.get("center", {}).get("lon")
        key = (name.lower(), round(lat or 0, 3), round(lon or 0, 3))
        if key in seen: continue
        seen.add(key)
        addr = ", ".join(x for x in [t.get("addr:housenumber"), t.get("addr:street"), t.get("addr:suburb") or t.get("addr:neighbourhood"), t.get("addr:city")] if x)
        row = {
            "niche": niche, "name": name,
            "category": t.get("amenity") or t.get("shop") or t.get("leisure") or t.get("office") or t.get("healthcare") or "",
            "phone": phone_of(t), "email": t.get("email") or t.get("contact:email", ""),
            "website": t.get("website") or t.get("contact:website", ""),
            "opening_hours": t.get("opening_hours", ""), "address": addr,
            "lat": lat, "lon": lon,
            "maps_link": f"https://www.openstreetmap.org/{el['type']}/{el['id']}",
            "osm_id": f"{el['type']}/{el['id']}",
            "kind": t.get("hairdresser") or t.get("beauty") or t.get("female") and "ladies" or "",
            "late": is_late(t.get("opening_hours", "")),
        }
        row["score"] = score(row)
        row["why_botbhai"] = PITCH[niche] + ("" if row["website"] else " No website listed — offer a free one-page site with the bot.")
        row["status"] = "new"; row["notes"] = ""
        rows.append(row)
    rows.sort(key=lambda r: -r["score"])
    return rows

COLS = ["score", "name", "category", "phone", "email", "website", "opening_hours", "address", "maps_link", "why_botbhai", "status", "notes", "lat", "lon", "osm_id", "niche", "kind", "late"]  # new columns go last: the sheet formulas use fixed letters

def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig opens cleanly in Excel
        w = csv.DictWriter(f, fieldnames=COLS, extrasaction="ignore"); w.writeheader(); w.writerows(rows)

def write_xlsx(path, by_niche, city):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import ColorScaleRule
    wb = Workbook(); summary = wb.active; summary.title = "Summary"
    head = PatternFill("solid", fgColor="128C7E"); white = Font(bold=True, color="FFFFFF")
    summary.append([f"BotBhai lead list — {city}"]); summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Data © OpenStreetMap contributors (ODbL). Verify before contacting; respect DND and never spam."])
    summary.append([]); summary.append(["Niche", "Businesses", "With phone", "No website", "Hot prospects (score ≥ 75)"])
    for c in summary[4]: c.fill, c.font = head, white
    for niche, rows in by_niche.items():
        n = len(rows) + 1
        summary.append([niche, f"=COUNTA('{niche}'!B2:B{n})", f"=COUNTIF('{niche}'!D2:D{n},\"?*\")", f"=COUNTBLANK('{niche}'!F2:F{n})", f"=COUNTIF('{niche}'!A2:A{n},\">=75\")"])
    for col, w in zip("ABCDE", (16, 14, 14, 14, 26)): summary.column_dimensions[col].width = w
    widths = {"score": 7, "name": 34, "category": 14, "phone": 18, "email": 24, "website": 30, "opening_hours": 22, "address": 40, "maps_link": 22, "why_botbhai": 60, "status": 12, "notes": 30}
    for niche, rows in by_niche.items():
        ws = wb.create_sheet(niche); ws.append(COLS)
        for c in ws[1]: c.fill, c.font = head, white
        for r in rows: ws.append([r.get(k, "") for k in COLS])
        for i, k in enumerate(COLS, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(k, 10)
        ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
        n = max(len(rows) + 1, 2)
        ws.conditional_formatting.add(f"A2:A{n}", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=50, mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"))
        dv = DataValidation(type="list", formula1='"new,contacted,demo shown,pilot,won,lost,do not contact"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"K2:K{n + 500}")
        for row in ws.iter_rows(min_row=2, min_col=10, max_col=10):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--city", default="Hyderabad")
    ap.add_argument("--niche", nargs="+", default=["salon", "gym", "realestate", "examkits"], choices=list(NICHE_TAGS))
    ap.add_argument("--out", default="leads")
    ap.add_argument("--near", help="only a walking cluster around 'lat,lon' (e.g. Ameerpet metro 17.4375,78.4482)")
    ap.add_argument("--radius", type=int, default=1200, help="cluster radius in metres (with --near)")
    ap.add_argument("--from-json", help="offline mode: build files from a saved Overpass JSON (testing)")
    a = ap.parse_args()
    by_niche = {}
    for n in a.niche:
        print(f"→ {n} in {a.city} …", file=sys.stderr)
        by_niche[n] = collect(a.city, n, parse_near(a.near) if a.near else None, a.radius)
        print(f"  {len(by_niche[n])} businesses, {sum(1 for r in by_niche[n] if r['phone'])} with phone", file=sys.stderr)
        time.sleep(5)  # be polite to the free public server
    allrows = [r for rows in by_niche.values() for r in rows]
    write_csv(a.out + ".csv", allrows); write_xlsx(a.out + ".xlsx", by_niche, a.city)
    print(f"Saved {len(allrows)} leads → {a.out}.csv, {a.out}.xlsx")

if __name__ == "__main__":
    main()
