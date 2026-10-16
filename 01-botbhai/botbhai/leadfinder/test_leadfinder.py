import sys, json; import os; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import leadfinder as lf
fake = {"elements": [
 {"type":"node","id":1,"lat":17.49,"lon":78.39,"tags":{"shop":"hairdresser","name":"Test Clinic A","phone":"+91 98480 00000","addr:street":"KPHB Road"}},
 {"type":"way","id":2,"center":{"lat":17.44,"lon":78.35},"tags":{"amenity":"dentist","name":"Smile Dental","website":"https://x.in","opening_hours":"Mo-Sa 10:00-20:00"}},
 {"type":"node","id":3,"lat":17.49,"lon":78.39,"tags":{"shop":"hairdresser","name":"Test Clinic A","phone":"+91 98480 00000"}},
 {"type":"node","id":4,"lat":17.4,"lon":78.4,"tags":{"amenity":"clinic"}}]}
lf.overpass = lambda q: fake
q = lf.build_query("Hyderabad", lf.NICHE_TAGS["salon"]); assert 'area["name"="Hyderabad"]' in q
rows = lf.collect("Hyderabad", "salon")
assert len(rows) == 2 and rows[0]["name"] == "Test Clinic A" and rows[0]["score"] == 95, rows
import tempfile; out = os.path.join(tempfile.mkdtemp(), "lf_test")
lf.write_csv(out+".csv", rows); lf.write_xlsx(out+".xlsx", {"salon": rows}, "Hyderabad")
from openpyxl import load_workbook; wb = load_workbook(out+".xlsx"); print(wb.sheetnames, wb["Summary"]["B5"].value, [c.value for c in wb["salon"][2]][:4])
print("leadfinder offline test passed")
# --- A08: one walking cluster (Overpass "around" filter), salon type kept for the pitch
q2 = lf.build_query("Hyderabad", lf.NICHE_TAGS["salon"], near=(17.4375, 78.4482), radius=1200)
assert 'nwr(around:1200,17.4375,78.4482)[shop~"^(hairdresser|beauty|cosmetics)$"];' in q2 and "area" not in q2, q2
fake["elements"].append({"type":"node","id":5,"lat":17.438,"lon":78.449,"tags":{"shop":"hairdresser","hairdresser":"unisex","name":"Style Unisex","opening_hours":"Mo-Su 10:00-22:00"}})
rows2 = lf.collect("Hyderabad", "salon", near=(17.4375, 78.4482), radius=1200)
r5 = next(r for r in rows2 if r["name"] == "Style Unisex"); assert r5["kind"] == "unisex" and r5["late"] is True, r5
assert next(r for r in rows2 if r["name"] == "Test Clinic A")["late"] is False
for bad in ["17.4,78.4", "abc", "95,78", "17.4,78.4,1"]:
    try: lf.parse_near(bad) if bad != "17.4,78.4" else None
    except SystemExit: continue
    assert bad == "17.4,78.4", bad
assert lf.parse_near("17.4375, 78.4482") == (17.4375, 78.4482)
print("leadfinder cluster test passed")
assert lf.COLS[0] == "score" and lf.COLS[3] == "phone" and lf.COLS[5] == "website" and lf.COLS[10] == "status", "xlsx formulas use fixed column letters"
