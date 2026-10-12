import sys, json; import os; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import leadfinder as lf
fake = {"elements": [
 {"type":"node","id":1,"lat":17.49,"lon":78.39,"tags":{"shop":"hairdresser","name":"Test Clinic A","phone":"+91 98480 00000","addr:street":"KPHB Road"}},
 {"type":"way","id":2,"center":{"lat":17.44,"lon":78.35},"tags":{"amenity":"dentist","name":"Smile Dental","website":"https://x.in","opening_hours":"Mo-Sa 10:00-20:00"}},
 {"type":"node","id":3,"lat":17.49,"lon":78.39,"tags":{"shop":"hairdresser","name":"Test Clinic A","phone":"+91 98480 00000"}},
 {"type":"node","id":4,"lat":17.4,"lon":78.4,"tags":{"amenity":"clinic"}}]}
lf.overpass = lambda q: fake
q = lf.build_query("Hyderabad", lf.NICHE_TAGS["salon"]); assert 'area["name"="Hyderabad"]' in q
rows = lf.collect("Hyderabad", "salon")
assert len(rows) == 2 and rows[0]["name"] == "Test Clinic A" and rows[0]["score"] == 95, rows
import tempfile; out = os.path.join(tempfile.mkdtemp(), "lf_test")
lf.write_csv(out+".csv", rows); lf.write_xlsx(out+".xlsx", {"salon": rows}, "Hyderabad")
from openpyxl import load_workbook; wb = load_workbook(out+".xlsx"); print(wb.sheetnames, wb["Summary"]["B5"].value, [c.value for c in wb["salon"][2]][:4])
print("leadfinder offline test passed")
