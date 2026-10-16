#!/usr/bin/env python3
"""A19 checks: build Lite / Kit / Pro with TODAY() pinned to 16 Oct 2026, fill realistic inputs, recalculate in
LibreOffice, and check the countdown, revision dates, streak, week totals, weights, mock analysis, tier contents,
no error cells, and Google-Sheets-portable functions only. Run: python3 03-kits/gate-2027/test_kit.py"""
import datetime as dt, os, pathlib, re, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook
HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build_kit as K  # noqa: E402

TODAY = dt.date(2026, 10, 16)
PORTABLE = {"IF", "AND", "OR", "MAX", "MIN", "SUM", "SUMIFS", "COUNTIF", "COUNTIFS", "COUNTA", "COUNT", "INDEX", "MATCH", "CHOOSE", "TODAY", "DATE",
            "ROUND", "AVERAGE", "IFERROR", "ISNUMBER", "N", "WEEKDAY"}


def recalc(src, tmp):
    out = pathlib.Path(tmp) / "recalc"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/lo", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return load_workbook(out / pathlib.Path(src).name, data_only=True)


def d(v):
    return v.date() if hasattr(v, "date") else v


def main():
    if not shutil.which("soffice"):
        if os.environ.get("CI"): sys.exit("LibreOffice missing on CI: the GATE kit was NOT checked")
        print("(GATE kit check skipped: needs LibreOffice)"); return
    n_topics = sum(len(t) for _, _, t in K.SUBJECTS)
    with tempfile.TemporaryDirectory() as tmp:
        files = {t: K.build(t, pathlib.Path(tmp) / f, TODAY) for t, f in K.FILES.items()}
        # functions used must exist in Google Sheets
        for t, f in files.items():
            wb = load_workbook(f)
            used = {m for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=") for m in re.findall(r"([A-Z][A-Z0-9.]*)\(", c.value)}
            assert used <= PORTABLE, f"{t}: non-portable functions {used - PORTABLE}"
        print("  ✓ all three files use only Google-Sheets-portable functions")
        # fill the Pro file like a real student
        wb = load_workbook(files["pro"]); sy = wb["Syllabus"]
        sy["C5"], sy["D5"], sy["E5"] = "Done", dt.datetime(2026, 10, 10), 0      # → revise 11 Oct: due
        sy["C6"], sy["D6"], sy["E6"] = "Done", dt.datetime(2026, 10, 10), 2      # → revise 17 Oct: later
        sy["C7"], sy["D7"], sy["E7"] = "Done", dt.datetime(2026, 10, 1), 4       # → all 4 revisions done: off the list
        sy["C8"] = "Learning"
        lg = wb["Daily log"]
        for r in range(6, 127):  # A6 = exam − 120 days = 9 Oct 2026
            day = dt.date(2027, 2, 6) - dt.timedelta(days=126 - r)
            if dt.date(2026, 10, 12) <= day <= dt.date(2026, 10, 15): lg[f"B{r}"] = 3
            if day == dt.date(2026, 10, 9): lg[f"B{r}"] = 2
        mk = wb["Mocks"]; mk["B6"], mk["C6"] = dt.datetime(2026, 10, 11), 41; mk["B7"], mk["C7"] = dt.datetime(2026, 10, 14), 47
        rows = {r[0].value: r[0].row for r in mk.iter_rows(min_row=29, max_row=40) if r[0].value}
        mk[f"B{rows['Computer Networks']}"], mk[f"C{rows['Computer Networks']}"] = 2, 5
        mk[f"B{rows['Algorithms']}"], mk[f"C{rows['Algorithms']}"] = 3, 1
        wb.save(files["pro"])
        R = {t: recalc(f, tmp) for t, f in files.items()}
    # no error cells anywhere
    for t, wb in R.items():
        errs = [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:"))]
        assert not errs, f"{t}: error cells {errs[:5]}"
    print("  ✓ no error cells in Lite, Kit or Pro after recalculation")
    p = R["pro"]; pl = p["Plan"]
    assert pl["B6"].value == (dt.date(2027, 2, 6) - TODAY).days == 113, pl["B6"].value
    assert pl["B7"].value == f"3 of {n_topics}", pl["B7"].value
    print(f"  ✓ countdown: 113 days to 6 Feb 2027; 3 of {n_topics} topics done")
    sy = p["Syllabus"]
    assert d(sy["F5"].value) == dt.date(2026, 10, 11) and sy["G5"].value == "Revise", (sy["F5"].value, sy["G5"].value)
    assert d(sy["F6"].value) == dt.date(2026, 10, 17) and sy["G6"].value == "Later", (sy["F6"].value, sy["G6"].value)
    assert sy["F7"].value in ("", None) and sy["F8"].value in ("", None)
    assert pl["B12"].value == 1, pl["B12"].value
    print("  ✓ spaced revision: +1 day (due), +7 days after 2 revisions (later), off the list after 4; 1 due today")
    assert p["Daily log"]["F3"].value == 4, p["Daily log"]["F3"].value
    assert pl["B11"].value == 4 and p["Daily log"]["B3"].value == 14
    wk = p["Week"]
    assert d(wk["B4"].value) == dt.date(2026, 10, 12) and wk["B5"].value == 12 and wk["B6"].value == 35, (wk["B4"].value, wk["B5"].value)
    print("  ✓ streak 4 (today not logged yet counts up to yesterday); 14 h total; this week Mon 12 Oct: 12 of 35 h")
    so = p["Sources"]; w = {so[f"A{r}"].value: so[f"E{r}"].value for r in range(6, 17)}
    assert w["Engineering Mathematics"] == 11.7 and w["Computer Organization and Architecture"] == 9.7 and w["General Aptitude"] == 15, w
    assert sy["H5"].value == 11.7
    m = p["Mocks"]
    assert m["F4"].value == 47 and m["F5"].value == 44 and m["F6"].value == 44
    leak = next(m[f"B{r}"].value for r in range(30, 50) if m[f"A{r}"].value == "Biggest leak")
    assert leak == "Computer Networks", leak
    fix = {m[f"A{r}"].value: m[f"F{r}"].value for r in range(29, 40)}
    assert fix["Computer Networks"] == "Re-learn" and fix["Algorithms"] == "Slow down" and fix["Databases"] in ("", None)
    print("  ✓ weights from GATE 2023–25 (Maths 11.7, COA 9.7, GA 15); mocks: best 47, avg 44; biggest leak Computer Networks (concept → re-learn)")
    names = {t: R[t].sheetnames for t in R}
    assert "Daily log" not in names["lite"] and "Mocks" not in names["lite"] and "Mocks" not in names["kit"] and "Daily log" in names["kit"]
    assert "తెలుగు · हिंदी" in names["pro"] and "తెలుగు · हिंदी" not in names["kit"]
    assert R["kit"]["Syllabus"]["H4"].value is None and R["pro"]["Syllabus"]["H4"].value == "Subject weight (avg marks)"
    assert R["lite"]["Syllabus"]["F5"].value is None
    print("  ✓ tiers: Lite = countdown + checklist; Kit + log, revision, week; Pro + weights, mocks, Telugu/Hindi")
    print("GATE kit tests passed")


if __name__ == "__main__":
    main()
