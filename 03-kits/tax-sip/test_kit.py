#!/usr/bin/env python3
"""A20 check: build the Tax + SIP kit, recalculate the real .xlsx in LibreOffice (headless), and compare with the
calculator site's tested engine (05-deals/site/src/lib/fin.ts, run with Node's TypeScript stripping):
  · the 3 official fixtures (₹12.75L / ₹13L / ₹25L) PASS on the in-sheet Self-check tab
  · 150 random salaries × deductions × age bands give exactly fin.ts compareRegimes() totals
  · HRA, SIP (no step-up) and EMI tabs equal fin.ts hra(), sip(), emi()
Needs: LibreOffice (soffice), Node ≥ 22.6, openpyxl.   Run: python3 03-kits/tax-sip/test_kit.py"""
import json, os, pathlib, random, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(HERE))
import build_kit as K  # noqa: E402

AGES = {"Below 60": "below60", "60 to 80": "60to80", "80 and above": "above80"}


def fin(js):
    code = f'import("{(ROOT / "05-deals/site/src/lib/fin.ts").as_posix()}").then((m) => {{ {js} }})'
    return json.loads(subprocess.check_output(["node", "--experimental-strip-types", "--no-warnings", "-e", code], text=True))


def recalc(src, tmp):
    out = pathlib.Path(tmp) / "recalc"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/lo-profile", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return load_workbook(out / pathlib.Path(src).name, data_only=True)


def main():
    if not shutil.which("soffice"):
        if os.environ.get("CI"): sys.exit("LibreOffice missing on CI: the kit was NOT checked")
        print("(tax kit check skipped: needs LibreOffice)"); return
    rnd = random.Random(2627)
    cases = [(rnd.randrange(200000, 5000000, 1000), rnd.choice([0, 50000, 150000, rnd.randrange(0, 450000, 500)]), rnd.choice(list(AGES))) for _ in range(150)]
    cases += [(1200000, 0, "Below 60"), (1275000, 0, "Below 60"), (1275010, 0, "Below 60"), (1287500, 0, "Below 60"), (550000, 0, "Below 60"), (550000, 0, "80 and above"), (5000000, 0, "Below 60")]
    want = fin("console.log(JSON.stringify(" + json.dumps([[g, d, AGES[a]] for g, d, a in cases]) + ".map(([g,d,a]) => { const c = m.compareRegimes(g,d,a); return [c.new.total, c.old.total]; })))")
    extra = [[g, d, a, n, o] for (g, d, a), (n, o) in zip(cases, want)]
    with tempfile.TemporaryDirectory() as tmp:
        src = pathlib.Path(tmp) / "kit.xlsx"; K.build(src, extra)
        # set the HRA / SIP / Loan inputs to known cases for the cross-checks
        wb = load_workbook(src)
        wb["HRA"]["B4"], wb["HRA"]["B5"], wb["HRA"]["B6"], wb["HRA"]["B7"] = 540000, 216000, 300000, "Other city"
        sp = wb["SIP goal"]; sp["B4"], sp["B5"], sp["B6"], sp["B7"], sp["B8"], sp["B9"] = 1000000, 12, 0, 0.11, 0, 0
        ln = wb["Loan"]; ln["B4"], ln["B5"], ln["B6"], ln["B7"] = 2500000, 0.09, 15, 0
        wb.save(src)
        r = recalc(src, tmp)
    s = r["Self-check"]
    rows = [[s.cell(i, c).value for c in range(1, 9)] for i in range(5, 5 + 3 + len(extra))]
    fx = rows[:3]
    assert all(x[7] == "PASS" for x in fx), f"official fixtures fail in the sheet: {fx}"
    print("  ✓ in-sheet Self-check: ₹12.75L / ₹13L / ₹25L → 0 / 26,000 / 3,19,800 new; 1,40,400 / 1,48,200 / 5,22,600 old: PASS")
    bad = [x for x in rows[3:] if x[7] != "PASS"]
    assert not bad, f"{len(bad)} of {len(extra)} cases differ from fin.ts, e.g. {bad[:3]}"
    print(f"  ✓ {len(extra)} random + edge cases (₹2L–₹50L, every age band, around the ₹12L rebate edge) equal the site's tax engine exactly")
    h = fin("const x = m.hra(540000, 216000, 300000, 'Other'); console.log(JSON.stringify(x.exempt))")
    assert r["HRA"]["B12"].value == h, (r["HRA"]["B12"].value, h)
    fv = fin("console.log(JSON.stringify(m.sip(1, 0.11, 12).value === undefined ? 0 : m.sip(100000, 0.11, 12).value / 100000))")
    assert abs(r["SIP goal"]["B14"].value - fv) < 0.01, (r["SIP goal"]["B14"].value, fv)
    e = fin("console.log(JSON.stringify(m.emi(2500000, 0.09, 15)))")
    assert r["Loan"]["B9"].value == e["emi"] and abs(r["Loan"]["B10"].value - e["totalInterest"]) <= 1 and r["Loan"]["B12"].value == 0 and r["Loan"]["B13"].value == 0, (r["Loan"]["B9"].value, r["Loan"]["B10"].value, e)
    print(f"  ✓ HRA ₹{h:,} · SIP factor {fv:.2f} · EMI ₹{e['emi']:,} (interest ₹{e['totalInterest']:,}) match fin.ts; no prepayment = no savings")
    for name in ["Start", "Tax", "HRA", "SIP goal", "Loan", "Self-check", "Rules"]:
        assert name in r.sheetnames, name
    print("tax kit tests passed")


if __name__ == "__main__":
    main()
